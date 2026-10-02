"""
Spare Parts Excel importer — a Spare Part Master/Sourcing Catalogue
import, never a stock import (no SparePartMovement is ever created here).

_scan_workbook() is the single source of truth for both the dry-run
preview and the real execute, so they can never disagree. Dedup identity
is normalized (name, model, size) — MACHINE is treated as a tag attached
via the many-to-many SparePartMachine association, not part of a spare's
identity, so the same physical spare reappearing against a different
machine column value attaches a new machine rather than creating a
duplicate master record.

Re-running the same workbook is idempotent via
(source_file, sheet_name, source_row) in SparePartImportRow — a source
row already recorded there is skipped outright, regardless of content.
"""
import hashlib
import json
import re
from decimal import Decimal, InvalidOperation

import openpyxl

from webapp.extensions import db
from webapp.models.spare_part import SparePart
from webapp.models.spare_part_import_row import SparePartImportRow
from webapp.services import spare_part_service as sp_service
from webapp.services.audit_service import record_audit

HEADER_ALIASES = {
    "#": "num", "spare part": "name", "model": "model", "size": "size",
    "machine": "machine", "name": "supplier_name", "building": "building",
    "contacts": "contacts", "price": "price",
}
# At least this many of these fields being the literal number zero marks a
# row as corrupted fill-down garbage rather than real data (e.g. the real
# workbook's row 23 — SIZE/MACHINE/NAME/BUILDING/CONTACTS/PRICE all 0.0).
_DEGENERATE_FIELDS = ("model", "size", "machine", "supplier_name", "building", "contacts", "price")
_DEGENERATE_THRESHOLD = 5


class SparePartImportError(ValueError):
    pass


def _norm_cell(v):
    if v is None:
        return None
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def _is_zeroish(v):
    if v is None:
        return False
    if isinstance(v, (int, float)):
        return v == 0
    return str(v).strip() in ("0", "0.0")


def parse_price(raw):
    """Returns a normalized Decimal only when unambiguous; else None. Never
    guesses on ranges, multi-quotes ("17k/15k"), or arithmetic text."""
    if raw is None:
        return None
    text = str(raw).strip().lower()
    if not text:
        return None
    if any(ch in text for ch in ("/", "+", "=", "-")):
        return None
    text = re.sub(r"\s+", "", text)
    text = text.replace(",", "")
    m = re.fullmatch(r"(\d+(?:\.\d+)?)(shs|ugx)?", text)
    if m:
        try:
            return Decimal(m.group(1))
        except InvalidOperation:
            return None
    m = re.fullmatch(r"(\d+(?:\.\d+)?)(k|m)", text)
    if m:
        try:
            value = Decimal(m.group(1))
        except InvalidOperation:
            return None
        return value * (Decimal(1000) if m.group(2) == "k" else Decimal(1000000))
    return None


def _find_header_row(ws, max_scan=5):
    for r in range(1, min(max_scan, ws.max_row) + 1):
        values = [str(ws.cell(row=r, column=c).value or "").strip().lower() for c in range(1, ws.max_column + 1)]
        if "spare part" in values:
            return r, values
    return None, None


def _row_fields(ws, header_map, row_num):
    raw_all = {}
    fields = {}
    for col in range(1, ws.max_column + 1):
        value = ws.cell(row=row_num, column=col).value
        if value in (None, ""):
            continue
        raw_all[col] = value
        key = header_map.get(col)
        if key:
            fields[key] = _norm_cell(value)
    return fields, raw_all


def _classify_row(fields):
    name = (fields.get("name") or "").strip()
    if not name:
        return "skip_blank", None
    zeroish_count = sum(1 for k in _DEGENERATE_FIELDS if _is_zeroish(fields.get(k)))
    if zeroish_count >= _DEGENERATE_THRESHOLD:
        return "ambiguous_invalid", "degenerate row (fill-down artifact — multiple fields are literal 0)"
    machine_raw = fields.get("machine")
    if machine_raw and sp_service.is_ambiguous_machine_text(machine_raw):
        return "ambiguous_machine", f"compound/ambiguous machine text: {machine_raw!r}"
    return "ok", None


def _identity_key(fields):
    return (
        sp_service.normalize_key(fields.get("name")),
        sp_service.normalize_key(fields.get("model")),
        sp_service.normalize_key(fields.get("size")),
    )


def _iter_data_rows(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        if ws.max_row <= 1:
            continue
        header_row, values = _find_header_row(ws)
        if header_row is None:
            continue
        header_map = {}
        for idx, label in enumerate(values, start=1):
            key = HEADER_ALIASES.get(label)
            if key:
                header_map[idx] = key
        for row_num in range(header_row + 1, ws.max_row + 1):
            fields, raw_all = _row_fields(ws, header_map, row_num)
            if not raw_all:
                continue
            yield sheet_name, row_num, fields, raw_all


def file_hash(path):
    h = hashlib.sha256()
    with open(path, "rb") as f:
        h.update(f.read())
    return h.hexdigest()[:16]


def _scan_workbook(path, persist=False, import_batch_id=None):
    """
    Single source of truth for preview and execute. persist=False never
    writes to the database (read-only classification against current DB
    state); persist=True actually creates SparePart/Machine/Supplier rows
    and SparePartImportRow provenance, skipping any source row already
    recorded in SparePartImportRow.
    """
    source_file = str(path)
    report = {
        "rows_detected": 0,
        "sheets": [],
        "new_spares": [],
        "attached_to_existing": [],
        "ambiguous_invalid": [],
        "ambiguous_machine": [],
        "skipped_already_imported": [],
        "machines_seen": set(),
        "suppliers_seen": set(),
    }
    seen_sheets = set()
    # In-batch identity map so two new rows within the SAME run that share
    # an identity key attach to the same new SparePart instead of each
    # creating one (mirrors how a rerun would resolve against the DB).
    batch_identity = {}

    for sheet_name, row_num, fields, raw_all in _iter_data_rows(path):
        seen_sheets.add(sheet_name)
        report["rows_detected"] += 1

        already = SparePartImportRow.query.filter_by(
            source_file=source_file, sheet_name=sheet_name, source_row=row_num,
        ).first()
        if already is not None:
            report["skipped_already_imported"].append({"sheet": sheet_name, "row": row_num})
            continue

        kind, reason = _classify_row(fields)
        if kind == "skip_blank":
            continue
        if kind == "ambiguous_invalid":
            report["ambiguous_invalid"].append({"sheet": sheet_name, "row": row_num, "reason": reason})
            continue

        if fields.get("machine"):
            report["machines_seen"].add(fields["machine"])
        if fields.get("supplier_name"):
            report["suppliers_seen"].add(fields["supplier_name"])
        if kind == "ambiguous_machine":
            report["ambiguous_machine"].append({"sheet": sheet_name, "row": row_num, "reason": reason})
            # Still imported — the ambiguity is about which machine(s) it
            # means, not about whether the spare itself is real.

        key = _identity_key(fields)
        existing_spare = batch_identity.get(key)
        if existing_spare is None:
            existing_spare = (
                SparePart.query
                .filter(db.func.lower(SparePart.name) == key[0])
                .filter(db.func.lower(db.func.coalesce(SparePart.model, "")) == key[1])
                .filter(db.func.lower(db.func.coalesce(SparePart.size, "")) == key[2])
                .first()
            )

        classification = "new_spares" if existing_spare is None else "attached_to_existing"

        if not persist:
            # Mark this identity as "seen" even in a dry run — without
            # this, a later row in the SAME preview pass that duplicates
            # an earlier one (e.g. Sheet2 repeating a Sheet1 row) would
            # look "new" too, since nothing gets committed mid-preview to
            # find via the DB query above. This keeps the dry run an
            # accurate preview of what execute_import will actually do.
            if existing_spare is None:
                batch_identity[key] = True
            report[classification].append({
                "sheet": sheet_name, "row": row_num, "name": fields.get("name"),
                "model": fields.get("model"), "size": fields.get("size"), "machine": fields.get("machine"),
            })
            continue

        if existing_spare is None:
            existing_spare = sp_service.create_spare_part(
                name=fields["name"], model=fields.get("model"), size=fields.get("size"), notes=None,
            )
            batch_identity[key] = existing_spare
            report["new_spares"].append({"sheet": sheet_name, "row": row_num, "spare_part_id": existing_spare.id})
        else:
            batch_identity[key] = existing_spare
            report["attached_to_existing"].append(
                {"sheet": sheet_name, "row": row_num, "spare_part_id": existing_spare.id}
            )

        machine = sp_service.normalize_machine(fields.get("machine"))
        if machine is not None:
            from webapp.models.spare_part import SparePartMachine
            link = SparePartMachine.query.filter_by(
                spare_part_id=existing_spare.id, machine_id=machine.id,
            ).first()
            if link is None:
                db.session.add(SparePartMachine(spare_part_id=existing_spare.id, machine_id=machine.id))

        if fields.get("supplier_name"):
            supplier = sp_service.find_or_create_supplier(
                fields.get("supplier_name"), building=fields.get("building"), contacts=fields.get("contacts"),
            )
            raw_price = fields.get("price")
            from webapp.models.spare_part import SparePartSupplier
            existing_quote = SparePartSupplier.query.filter_by(
                spare_part_id=existing_spare.id, supplier_id=supplier.id, raw_price_text=_norm_cell(raw_price),
            ).first()
            if existing_quote is None:
                db.session.add(SparePartSupplier(
                    spare_part_id=existing_spare.id, supplier_id=supplier.id,
                    raw_price_text=_norm_cell(raw_price), normalized_price=parse_price(raw_price),
                    notes=None,
                ))

        db.session.flush()
        db.session.add(SparePartImportRow(
            source_file=source_file, sheet_name=sheet_name, source_row=row_num,
            raw_json=json.dumps({str(c): _norm_cell(v) for c, v in raw_all.items()}),
            spare_part_id=existing_spare.id, import_batch_id=import_batch_id or "manual",
        ))

    report["sheets"] = sorted(seen_sheets)
    report["machines_seen"] = sorted(report["machines_seen"])
    report["suppliers_seen"] = sorted(report["suppliers_seen"])
    return report


def preview_import(path):
    report = _scan_workbook(path, persist=False)
    report["preview_token"] = file_hash(path)
    return report


def execute_import(path, actor, preview_token, import_batch_id=None):
    if preview_token != file_hash(path):
        raise SparePartImportError(
            "Preview token does not match this file's current contents — re-run the dry run first."
        )
    report = _scan_workbook(path, persist=True, import_batch_id=import_batch_id)
    record_audit(actor, "import_execute", "spare_part_import", entity_id=path, after={
        "rows_detected": report["rows_detected"],
        "new_spares": len(report["new_spares"]),
        "attached_to_existing": len(report["attached_to_existing"]),
        "ambiguous_invalid": len(report["ambiguous_invalid"]),
        "skipped_already_imported": len(report["skipped_already_imported"]),
    })
    return report
