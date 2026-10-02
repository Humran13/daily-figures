"""
Spare Parts Excel importer — tested against a small synthetic workbook
built in a tmp fixture (not the real "MV MACHINE SPARES.xlsx", which
stays out of the automated suite and out of the repo). Mirrors the real
workbook's structure: Sheet1 primary data, Sheet2 partially duplicating
Sheet1, Sheet3 blank, a machine-case-variant pair, and a corrupted
fill-down row.
"""
import openpyxl
import pytest

from webapp.models.spare_part import SparePart
from webapp.models.spare_part_import_row import SparePartImportRow
from webapp.services import spare_part_import_service as import_svc


@pytest.fixture
def workbook_path(tmp_path):
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "Sheet1"
    ws1.append(["MOTION VENTURES LTD MACHINE SPARES"])
    ws1.append(["#", "SPARE PART", "MODEL", "SIZE", "MACHINE", "NAME", "BUILDING", "CONTACTS", "PRICE"])
    ws1.append([1, "Temperature Controller", "Digital", "Medium", "tp", "Central Town", "Ham Building", 256704127039, "150,000shs"])
    ws1.append([2, "Bearings Hajj", "Mechanical", '9"', "TP", "Down Town", "Nakivubo", 256702460512, "100,000shs"])
    ws1.append([3, "V-Belt pix", "B52", 0, 0, 0, 0, 0, 0])  # corrupted fill-down row
    ws1.append([4, "Carbon Brush", "J204", "21g", "np/tp", "Central Town", "Ham Building", None, "25k"])
    ws1.append([5, "Allen Bolt M6", None, None, None, None, None, None, None])

    ws2 = wb.create_sheet("Sheet2")
    ws2.append(["MOTION VENTURES LTD MACHINE SPARES"])
    ws2.append(["#", "SPARE PART", "MODEL", "SIZE", "MACHINE", "NAME", "BUILDING", "CONTACTS", "PRICE"])
    ws2.append([1, "Temperature Controller", "Digital", "Medium", "tp", "Central Town", "Ham Building", 256704127039, "150,000shs"])  # duplicate of Sheet1 row 3
    ws2.append([2, "Unique Sheet2 Spare", "ModelX", "Large", "gn", "Other Town", "Other Building", None, "80k"])  # genuinely unique

    wb.create_sheet("Sheet3")  # blank

    path = tmp_path / "test_spares.xlsx"
    wb.save(path)
    return str(path)


# ---------- parsing / price ----------

def test_parse_price_unambiguous_forms():
    from decimal import Decimal
    assert import_svc.parse_price("150,000shs") == Decimal("150000")
    assert import_svc.parse_price("100000") == Decimal("100000")
    assert import_svc.parse_price("85k") == Decimal("85000")
    assert import_svc.parse_price("2m") == Decimal("2000000")


def test_parse_price_ambiguous_forms_return_none():
    assert import_svc.parse_price("17k/15k") is None
    assert import_svc.parse_price("6.9+4.44=11.4") is None
    assert import_svc.parse_price(None) is None


# ---------- workbook scan ----------

def test_workbook_can_be_parsed(app, workbook_path):
    with app.app_context():
        report = import_svc.preview_import(workbook_path)
        assert report["rows_detected"] > 0


def test_sheet1_rows_recognized(app, workbook_path):
    with app.app_context():
        report = import_svc.preview_import(workbook_path)
        assert "Sheet1" in report["sheets"]


def test_blank_sheet_ignored(app, workbook_path):
    with app.app_context():
        report = import_svc.preview_import(workbook_path)
        assert "Sheet3" not in report["sheets"]


def test_dry_run_makes_no_db_changes(app, workbook_path):
    with app.app_context():
        import_svc.preview_import(workbook_path)
        assert SparePart.query.count() == 0
        assert SparePartImportRow.query.count() == 0


def test_corrupted_fill_down_row_flagged_ambiguous_invalid(app, workbook_path):
    with app.app_context():
        report = import_svc.preview_import(workbook_path)
        assert any(r["row"] == 5 for r in report["ambiguous_invalid"])  # row 5 = the 0/0/0 row


def test_compound_machine_text_flagged_ambiguous(app, workbook_path):
    with app.app_context():
        report = import_svc.preview_import(workbook_path)
        assert any("np/tp" in r["reason"] for r in report["ambiguous_machine"])


# ---------- real import ----------

def test_real_import_works(app, workbook_path):
    with app.app_context():
        from webapp.extensions import db
        from webapp.models.user import ROLE_SUPER_ADMIN, User
        from werkzeug.security import generate_password_hash
        actor = User(username="importer", password_hash=generate_password_hash("x"), role=ROLE_SUPER_ADMIN)
        db.session.add(actor)
        db.session.commit()

        token = import_svc.file_hash(workbook_path)
        result = import_svc.execute_import(workbook_path, actor, token)
        db.session.commit()

        assert len(result["new_spares"]) > 0
        assert SparePart.query.filter_by(name="Temperature Controller").count() == 1


def test_sheet2_duplicate_not_duplicated(app, workbook_path):
    with app.app_context():
        from webapp.extensions import db
        from webapp.models.user import ROLE_SUPER_ADMIN, User
        from werkzeug.security import generate_password_hash
        actor = User(username="importer2", password_hash=generate_password_hash("x"), role=ROLE_SUPER_ADMIN)
        db.session.add(actor)
        db.session.commit()

        token = import_svc.file_hash(workbook_path)
        import_svc.execute_import(workbook_path, actor, token)
        db.session.commit()

        # Sheet2's row 3 is a content-duplicate of Sheet1's "Temperature
        # Controller" row — must not create a second SparePart.
        assert SparePart.query.filter_by(name="Temperature Controller").count() == 1


def test_sheet2_unique_row_imported(app, workbook_path):
    with app.app_context():
        from webapp.extensions import db
        from webapp.models.user import ROLE_SUPER_ADMIN, User
        from werkzeug.security import generate_password_hash
        actor = User(username="importer3", password_hash=generate_password_hash("x"), role=ROLE_SUPER_ADMIN)
        db.session.add(actor)
        db.session.commit()

        token = import_svc.file_hash(workbook_path)
        import_svc.execute_import(workbook_path, actor, token)
        db.session.commit()

        assert SparePart.query.filter_by(name="Unique Sheet2 Spare").count() == 1


def test_repeated_import_is_idempotent(app, workbook_path):
    with app.app_context():
        from webapp.extensions import db
        from webapp.models.user import ROLE_SUPER_ADMIN, User
        from werkzeug.security import generate_password_hash
        actor = User(username="importer4", password_hash=generate_password_hash("x"), role=ROLE_SUPER_ADMIN)
        db.session.add(actor)
        db.session.commit()

        token = import_svc.file_hash(workbook_path)
        import_svc.execute_import(workbook_path, actor, token)
        db.session.commit()
        count_after_first = SparePart.query.count()
        import_row_count_after_first = SparePartImportRow.query.count()

        result2 = import_svc.execute_import(workbook_path, actor, token)
        db.session.commit()

        assert SparePart.query.count() == count_after_first
        assert SparePartImportRow.query.count() == import_row_count_after_first
        assert len(result2["new_spares"]) == 0
        assert len(result2["skipped_already_imported"]) > 0


def test_machine_case_whitespace_aliases_normalize_safely(app, workbook_path):
    with app.app_context():
        from webapp.extensions import db
        from webapp.models.machine import Machine
        from webapp.models.user import ROLE_SUPER_ADMIN, User
        from werkzeug.security import generate_password_hash
        actor = User(username="importer5", password_hash=generate_password_hash("x"), role=ROLE_SUPER_ADMIN)
        db.session.add(actor)
        db.session.commit()

        token = import_svc.file_hash(workbook_path)
        import_svc.execute_import(workbook_path, actor, token)
        db.session.commit()

        # "tp" (row 3) and "TP" (row 4) must resolve to the same Machine.
        tp_machines = [m for m in Machine.query.all() if m.name.lower() == "tp"]
        assert len(tp_machines) == 1


def test_ambiguous_rows_are_flagged_not_merged(app, workbook_path):
    with app.app_context():
        from webapp.extensions import db
        from webapp.models.user import ROLE_SUPER_ADMIN, User
        from werkzeug.security import generate_password_hash
        actor = User(username="importer6", password_hash=generate_password_hash("x"), role=ROLE_SUPER_ADMIN)
        db.session.add(actor)
        db.session.commit()

        token = import_svc.file_hash(workbook_path)
        result = import_svc.execute_import(workbook_path, actor, token)
        db.session.commit()

        # The corrupted all-zero row is never persisted as a SparePart.
        assert SparePart.query.filter_by(name="V-Belt pix").count() == 0
        assert len(result["ambiguous_invalid"]) == 1


def test_raw_source_information_remains_traceable(app, workbook_path):
    with app.app_context():
        from webapp.extensions import db
        from webapp.models.user import ROLE_SUPER_ADMIN, User
        from werkzeug.security import generate_password_hash
        import json
        actor = User(username="importer7", password_hash=generate_password_hash("x"), role=ROLE_SUPER_ADMIN)
        db.session.add(actor)
        db.session.commit()

        token = import_svc.file_hash(workbook_path)
        import_svc.execute_import(workbook_path, actor, token)
        db.session.commit()

        row = SparePartImportRow.query.filter_by(sheet_name="Sheet1", source_row=3).first()
        assert row is not None
        raw = json.loads(row.raw_json)
        assert "Temperature Controller" in raw.values()


def test_import_never_creates_stock_quantities(app, workbook_path):
    with app.app_context():
        from webapp.extensions import db
        from webapp.models.spare_part_movement import SparePartMovement
        from webapp.models.user import ROLE_SUPER_ADMIN, User
        from werkzeug.security import generate_password_hash
        actor = User(username="importer8", password_hash=generate_password_hash("x"), role=ROLE_SUPER_ADMIN)
        db.session.add(actor)
        db.session.commit()

        token = import_svc.file_hash(workbook_path)
        import_svc.execute_import(workbook_path, actor, token)
        db.session.commit()

        assert SparePartMovement.query.count() == 0


def test_execute_rejects_stale_preview_token(app, workbook_path):
    with app.app_context():
        from webapp.models.user import ROLE_SUPER_ADMIN, User
        from werkzeug.security import generate_password_hash
        actor = User(username="importer9", password_hash=generate_password_hash("x"), role=ROLE_SUPER_ADMIN)
        from webapp.extensions import db
        db.session.add(actor)
        db.session.commit()

        with pytest.raises(import_svc.SparePartImportError):
            import_svc.execute_import(workbook_path, actor, "stale-token-0000")
